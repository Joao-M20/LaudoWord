import pandas as pd
from docx import Document

# Caminho do arquivo Excel
excel_path = 'C:\\Users\\BurgerKing\\Downloads\\LAUDO - SUCATA BK.xlsx'

# Caminho do arquivo Word
word_path = 'C:\\Users\\BurgerKing\\Documents\\LAUDO - SUCATA 116 Teste.docx'

# Caminho para salvar o Word preenchido
output_path = 'C:\\Users\\BurgerKing\\Documents\\preenchido.docx'

# Ler a planilha do Excel
df = pd.read_excel(excel_path, sheet_name='LAUDO TESTE')

# Filtrar os dados da coluna "N de serie"
numeros_serie = df['N de serie'].tolist()  # Converte a coluna para uma lista

# Abrir o documento Word
doc = Document(word_path)

# Variável para controlar o índice dos números de série
index = 0

# Substituir placeholders dentro de tabelas
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if "{{serial}}" in cell.text:  # Verifica se há um placeholder
                if index < len(numeros_serie):  # Verifica se há números de série disponíveis
                    cell.text = cell.text.replace("{{serial}}", str(numeros_serie[index]))
                    index += 1  # Avança para o próximo número de série

# Salvar o documento preenchido
doc.save(output_path)

print(f"Documento preenchido salvo em: {output_path}")
